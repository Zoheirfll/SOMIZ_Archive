"""Génération des fichiers d'export des registres : fonctions pures (titre,
entêtes, lignes) → octets, pour que Excel et PDF restent testables sans HTTP."""
from datetime import date, datetime
from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMULE_DEBUTS = ('=', '+', '-', '@')


def neutraliser(valeur):
    """Les champs libres (observation, nom) viennent d'utilisateurs : une
    cellule commençant par = + - @ serait interprétée comme une formule par
    Excel (injection). On la préfixe d'une apostrophe."""
    if isinstance(valeur, str) and valeur.startswith(FORMULE_DEBUTS):
        return "'" + valeur
    return valeur


def construire_xlsx(titre, entetes, lignes, colonnes_date=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append(entetes)
    for col, entete in enumerate(entetes, start=1):
        cellule = ws.cell(row=1, column=col)
        cellule.font = Font(bold=True)
        ws.column_dimensions[cellule.column_letter].width = max(len(str(entete)) + 2, 14)
    for ligne in lignes:
        ws.append([neutraliser(v) for v in ligne])
    for row in ws.iter_rows(min_row=2):
        for idx in colonnes_date:
            if isinstance(row[idx].value, (date, datetime)):
                row[idx].number_format = 'DD/MM/YYYY'
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _texte_pdf(valeur):
    """Texte d'une cellule PDF. Paragraph interprète du balisage : les champs
    libres sont échappés."""
    if isinstance(valeur, (date, datetime)):
        return valeur.strftime('%d/%m/%Y')
    texte = '' if valeur is None else str(valeur)
    return texte.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def construire_pdf(titre, entetes, lignes):
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
        title=titre,
    )
    styles = getSampleStyleSheet()
    cellule = styles['BodyText']
    cellule.fontSize = 7
    cellule.leading = 8.5
    donnees = [[Paragraph(f"<b>{_texte_pdf(e)}</b>", cellule) for e in entetes]]
    for ligne in lignes:
        donnees.append([Paragraph(_texte_pdf(v), cellule) for v in ligne])
    table = Table(donnees, repeatRows=1)
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    doc.build([Paragraph(_texte_pdf(titre), styles['Title']), Spacer(1, 8), table])
    return buf.getvalue()
